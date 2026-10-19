from datetime import date, datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook

from src.analysis import PULL_NOT, PULL_PARTIAL, PULL_YES, apply_filters, aggregate_stops, stopped_by_location
from src.export import cables_at_location, catra_list_xlsx
from src.progress import KIND_COMPLETE, KIND_PARTIAL, build_pull_events, daily_progress, progress_stats
from src.categories import extract_k_category, thickness_label
from src.dates import parse_flexible_date
from src.ids import join_key, normalize_kabel_id
from src.join import join_cable_data
from src.parsers import parse_catra_list, parse_status_list


def test_normalize_and_join_keys():
    assert normalize_kabel_id("0002 (E) K3") == "0002"
    assert normalize_kabel_id("0002,1 (E) K3") == "0002,1"
    assert normalize_kabel_id(2114.1) == "2114,1"
    assert join_key("0002 (E) K3") == "2"
    assert join_key("0002,1") == "2,1"
    assert join_key("5012.1") == join_key("5012,1") == "5012,1"
    assert join_key("05012.1") == "5012,1"

    from src.ids import cable_id_list_stats, parse_cable_id_list, select_cables_by_ids

    assert parse_cable_id_list("5012.1; 0006; 5012,1") == ["5012,1", "6"]
    assert parse_cable_id_list("5012.1\n0006") == ["5012,1", "6"]
    assert parse_cable_id_list("5012.1; 0006; 5012,1", dedupe=False) == ["5012,1", "6", "5012,1"]

    stats = cable_id_list_stats("5012.1; 0006; 5012,1; 0006")
    assert stats["raw_count"] == 4
    assert stats["unique_count"] == 2
    assert stats["duplicate_count"] == 2
    assert stats["unique"] == ["5012,1", "6"]
    assert set(stats["duplicates"]) == {"5012,1", "6"}

    sample = pd.DataFrame(
        {
            "Kabelnr": ["5012,1", "0006", "3602"],
            "join_key": ["5012,1", "6", "3602"],
            "Length_m": [10, 20, 30],
        }
    )
    matched, found, missing = select_cables_by_ids(sample, "5012.1; 9999; 0006")
    assert list(matched["join_key"]) == ["5012,1", "6"]
    assert found == ["5012,1", "6"]
    assert missing == ["9999"]
    # Order follows the request list, not the dataframe order
    matched2, _, _ = select_cables_by_ids(sample, "3602; 5012,1")
    assert list(matched2["join_key"]) == ["3602", "5012,1"]


def test_export_cache_key_includes_cable_selection():
    """Changing ID selection must not reuse a previous CaTra xlsx payload."""
    from openpyxl import load_workbook

    from src.export import catra_list_xlsx
    from src.ui import _cable_export_fingerprint

    a = pd.DataFrame(
        {
            "join_key": ["1632", "1633"],
            "Kabelnr": ["1632", "1633"],
            "_block": [
                [["1632", "A", None, 10], ["x"], ["y"], ["1E1 (E 1,0) 1E2"]],
                [["1633", "A", None, 20], ["x"], ["y"], ["2E1 (E 1,0) 2E2"]],
            ],
        }
    )
    b = pd.DataFrame(
        {
            "join_key": ["9999"],
            "Kabelnr": ["9999"],
            "_block": [[["9999", "A", None, 5], ["x"], ["y"], ["9E1"]]],
        }
    )
    assert _cable_export_fingerprint(a) != _cable_export_fingerprint(b)
    assert _cable_export_fingerprint(a) == ("1632", "1633")

    payload_a = catra_list_xlsx(a)
    payload_b = catra_list_xlsx(b)
    assert payload_a != payload_b
    sheet_a = load_workbook(BytesIO(payload_a)).active
    sheet_b = load_workbook(BytesIO(payload_b)).active
    assert "1632" in str(sheet_a["A5"].value)
    assert "9999" in str(sheet_b["A5"].value)
    assert "1632" not in str(sheet_b["A5"].value)


def test_dates_serial_and_string():
    assert parse_flexible_date("26.08.2026").strftime("%Y-%m-%d") == "2026-08-26"
    assert parse_flexible_date(46227).year == 2026
    assert pd.isna(parse_flexible_date(""))
    assert pd.isna(parse_flexible_date("nein"))
    assert pd.isna(parse_flexible_date(datetime(1677, 9, 22)))
    assert pd.isna(parse_flexible_date("01.01.1800"))

    ancient = pd.DataFrame(
        {
            "Kabelnr": ["1"],
            "Length_m": [100],
            "Gestoppt_1": [pd.Timestamp("1677-09-22")],
            "Teilgezogene_meter_1": [10],
            "verlegt": [pd.NaT],
        }
    )
    assert build_pull_events(ancient).empty


def test_categories_and_thickness():
    assert extract_k_category("0005 (E) K1") == "Cat1"
    assert extract_k_category("4010 (S) K5") == "Cat5"
    assert thickness_label("3x95") == "Thick"
    assert thickness_label("3G2,5") == "Thin"
    assert thickness_label("16x1,5") == "Thin"
    assert thickness_label("Simplex") == "Thin"


def _status_xlsx() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "Kabelnr",
            "Liste",
            "Zuschn.",
            "verlegt",
            "A.Gerät",
            "E.Gerät",
            "Kabeladern",
            "Startsegment_1",
            "Endsegment_1",
            "Gestoppt_1",
            "Teilgezogene meter_1",
        ]
    )
    ws.append(["0002", "0019", 128, "", "67 Converter", "103.6 CB", "Simplex", "3EDV204", "3EDH407", "31.08.2026", 91.2])
    ws.append(["0006", "0008", 31, "17.08.2026", "136.1 DG", "67 Converter", "3x95", "1E1206", "1EDV207.1", "", ""])
    ws.append(["0007", "0008", 31, "", "136.1 DG", "67 Converter", "3x95", "", "", "", ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _catra_xlsx() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.append(["CaTra - Kabelliste"])
    ws.append(["Objekt: 598"])
    ws.append(["Erstellt"])
    ws.append(["Kabel", "Artikel / Liste", None, "Länge", None, None, "Knoten", "A.län.", "Ort", "Gerät"])
    ws.append(["0002 (E) K3", "Ship Data Cat.7", "Simplex", 128, None, "A:", "1E2104", 6.5, "1D #20-28 SB", "67 Converter"])
    ws.append(["142", "0019", None, None, None, "E:", "3E4005", 14, "3D #159 PS", "103.6 CB"])
    ws.append(["L0019 Stolze", None, None, None, None, "Z:", "3E3103"])
    ws.append(["1E2104 (E 2,0) 1E2105"])
    ws.append(["0006 (E) K1", "Ship PWR/FLX SC 1kV", "3x95", 31, None, "A:", "1E1206", 4, "1D #12 PS", "136.1 DG"])
    ws.append(["6500", "0008", None, None, None, "E:", "1EDV207.1", 4, "1D #20-28 SB", "67 Converter"])
    ws.append(["Feld5", None, None, None, None, "Z:"])
    ws.append(["1E1206 1EDH219"])
    ws.append(["0007 (E) K1", "Ship PWR/FLX SC 1kV", "3x95", 31, None, "A:", "1E1206", 4, "1D #12 PS", "136.1 DG"])
    ws.append(["6500", "0008", None, None, None, "E:", "1EDV207.1", 4, "1D #20-28 SB", "67 Converter"])
    ws.append(["Feld5", None, None, None, None, "Z:"])
    ws.append(["1E1206 1EDH219"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_parse_join_filters_and_stops():
    status = parse_status_list(_status_xlsx())
    catra = parse_catra_list(_catra_xlsx())
    joined = join_cable_data(status, catra)
    assert set(joined["join_key"]) == {"2", "6", "7"}
    assert joined.loc[joined["join_key"] == "2", "Category"].iloc[0] == "Cat3"
    assert joined.loc[joined["join_key"] == "6", "Thickness"].iloc[0] == "Thick"
    # Liste from L#### note when present
    assert catra.loc[catra["join_key"] == "2", "Liste"].iloc[0] == "L0019"

    not_pulled = apply_filters(joined, pulled=[PULL_NOT])
    assert list(not_pulled["join_key"]) == ["7"]
    assert not_pulled["Startsegment_1"].fillna("").eq("").all()

    partial = apply_filters(joined, pulled=[PULL_PARTIAL])
    assert list(partial["join_key"]) == ["2"]

    both = apply_filters(joined, pulled=[PULL_NOT, PULL_PARTIAL])
    assert set(both["join_key"]) == {"2", "7"}

    laid = apply_filters(joined, pulled=[PULL_YES])
    assert list(laid["join_key"]) == ["6"]

    from src.analysis import (
        cable_numbers,
        cable_numbers_by_type,
        long_list_cables,
        pulled_cables,
        started_cables,
        total_length,
    )

    mixed = apply_filters(joined, pulled=[PULL_PARTIAL, PULL_YES])
    assert round(total_length(mixed), 1) == round(91.2 + 31, 1)
    assert round(total_length(apply_filters(joined, pulled=[PULL_PARTIAL])), 1) == 91.2

    assert cable_numbers(pulled_cables(joined)) == ["6"]
    assert cable_numbers(started_cables(joined)) == ["2"]

    sample = pd.DataFrame(
        {
            "Kabelnr": ["10", "11", "12", "13"],
            "Liste": ["001", "0001", "0019", "001"],
            "Length_m": [200, 151, 400, 149],
            "Kabeltyp": ["Ship PWR", "Ship Data", "Ship PWR", "Ship PWR"],
            "Kabeladern": ["3x95", "Simplex", "3x95", "3x95"],
        }
    )
    long_list = long_list_cables(sample, liste="001", min_m=150)
    assert set(long_list["Kabelnr"]) == {"10", "11"}
    grouped = cable_numbers_by_type(long_list)
    assert grouped["Ship PWR · Adern: 3x95"] == ["10"]
    assert grouped["Ship Data · Adern: Simplex"] == ["11"]

    device = apply_filters(joined, devices=["67 Converter"])
    assert len(device) == 3

    stops = stopped_by_location(joined)
    summary = aggregate_stops(stops)
    row = summary.loc[summary["Stop_location"] == "3EDH407"].iloc[0]
    assert row["Cable_count"] == 1
    assert round(row["Remaining_unpulled_m"], 1) == 36.8
    assert row["Total_cable_length_m"] == 128

    finished = joined.copy()
    finished.loc[finished["join_key"] == "2", "verlegt"] = pd.Timestamp("2026-09-01")
    finished.loc[finished["join_key"] == "2", "Pulled"] = True
    finished.loc[finished["join_key"] == "2", "Pull_status"] = PULL_YES
    assert stopped_by_location(finished).empty

    continued = joined.loc[joined["join_key"] == "2"].copy()
    continued["Gestoppt_2"] = pd.Timestamp("2026-09-02")
    continued["Startsegment_2"] = "3EDH407"
    continued["Endsegment_2"] = "3EDV300"
    continued["Teilgezogene_meter_2"] = 10
    later = stopped_by_location(continued)
    assert list(later["Stop_location"]) == ["3EDV300"]

    at_final = joined.loc[joined["join_key"] == "2"].copy()
    at_final["Endsegment_1"] = at_final["Knoten_E"]
    assert stopped_by_location(at_final).empty

    events = build_pull_events(joined)
    assert set(events["Kind"]) == {KIND_PARTIAL, KIND_COMPLETE}
    assert events.loc[events["Kind"].eq(KIND_PARTIAL), "Meters"].sum() == 91.2
    assert events.loc[events["Kind"].eq(KIND_COMPLETE), "Meters"].sum() == 31
    daily = daily_progress(events)
    stats = progress_stats(events, daily)
    assert round(stats["avg_m_per_day"], 2) == round((91.2 + 31) / stats["days"], 2)
    only_partial = daily_progress(events, [KIND_PARTIAL])
    assert only_partial["Total"].sum() == 91.2
    from src.progress import weekly_hours, weekly_progress

    weekly = weekly_progress(daily)
    assert round(float(weekly["Total"].sum()), 1) == round(91.2 + 31, 1)
    assert weekly["KW"].str.match(r"KW\d{2} \d{4}").all()
    week_36 = weekly.loc[weekly["KW"].eq("KW36 2026")]
    assert not week_36.empty
    assert round(float(week_36["Total"].iloc[0]), 1) == 91.2

    hrs = weekly_hours(
        pd.DataFrame(
            {
                "Date": pd.to_datetime(["2026-08-31", "2026-09-01", "2026-09-07"]),
                "Hours": [10.0, 20.0, 5.0],
            }
        )
    )
    assert round(float(hrs.loc[hrs["KW"].eq("KW36 2026"), "Hours"].iloc[0]), 1) == 30.0
    assert round(float(hrs.loc[hrs["KW"].eq("KW37 2026"), "Hours"].iloc[0]), 1) == 5.0

    loc_cables = cables_at_location(joined, stops, "3EDH407")
    assert list(loc_cables["join_key"]) == ["2"]
    assert loc_cables.iloc[0]["_block"][0][0] == "0002 (E) K3"

    from src.pathways import cables_stopped_at, stop_locations_series, unique_stop_stuetzen

    assert unique_stop_stuetzen(joined) == ["3EDH407"]
    assert stop_locations_series(joined).tolist()[0] == ["3EDH407"]
    at_stuetze = cables_stopped_at(joined, ["3EDH407"])
    assert list(at_stuetze["join_key"]) == ["2"]
    assert cables_stopped_at(joined, ["2EDV208"]).empty
    exported = catra_list_xlsx(loc_cables, list(catra.attrs.get("catra_header") or []))
    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(exported)).active
    assert str(sheet["A5"].value).startswith("0002")
    assert "1E2104" in str(sheet["A8"].value)


def test_catra_export_keeps_formula_like_text():
    from zipfile import ZipFile

    from openpyxl import load_workbook

    from src.export import catra_list_xlsx

    cables = pd.DataFrame(
        [
            {
                "join_key": "1",
                "Kabelnr": "0001",
                "_block": [
                    ["0001 (E) K3", "=A1+1", None, 10, None, "A:", "3E4003", 1, "x", "6 4D #152-1"],
                    ["note", None, None, None, None, "E:", "3E4004"],
                    ["=SUM(1,2)", None, None, None, None, "Z:"],
                    ["3EDH407.1 (E 0,0) 3E4003 (E 5,0) 4EDV400"],
                ],
            }
        ]
    )
    exported = catra_list_xlsx(cables)
    with ZipFile(BytesIO(exported)) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<f" not in xml
    sheet = load_workbook(BytesIO(exported)).active
    assert sheet["B5"].value == "=A1+1"
    assert sheet["A7"].value == "=SUM(1,2)"
    assert "3E4003" in str(sheet["A8"].value)


def test_pathway_wraps_in_export():
    from openpyxl import load_workbook

    from src.export import catra_list_xlsx, wrap_pathway

    long_path = (
        "3EDH407.1 (E 0,0) 3EDV407.1 (E 5,3) 3E4003 (E 5,0) 4EDV400 (D 0,0) "
        "4EDV400.1 (E 2,0) 4E2104 (E 3,1) 4E2105 (E 3,1) 4E4004"
    )
    wrapped = wrap_pathway(long_path)
    assert "\n" in wrapped
    assert "3EDH407.1" in wrapped.splitlines()[0]
    assert wrapped.replace("\n", " ") == long_path

    cables = pd.DataFrame(
        [
            {
                "join_key": "1",
                "Kabelnr": "0001",
                "Pfad": long_path,
                "_block": [
                    ["0001 (E) K3"],
                    ["note"],
                    ["z"],
                    [long_path],
                ],
            }
        ]
    )
    sheet = load_workbook(BytesIO(catra_list_xlsx(cables))).active
    assert "\n" in str(sheet["A8"].value)
    assert sheet["A8"].alignment.wrap_text is True


def test_kurzliste_one_row_with_path():
    from openpyxl import load_workbook

    from src.export import kurzliste_xlsx

    cables = pd.DataFrame(
        [
            {
                "join_key": "1",
                "Kabelnr": "0001",
                "Liste": "0019",
                "Artikel": "Ship Data",
                "Typ": "Simplex",
                "Length_m": 128,
                "Knoten_A": "1E2104",
                "Knoten_E": "3E4005",
                "Pfad": "1E2104 (E 2,0) 1E2105 (E 1,0) 3E4005",
            }
        ]
    )
    sheet = load_workbook(BytesIO(kurzliste_xlsx(cables))).active
    assert sheet["A1"].value == "Kabelnr"
    assert str(sheet["A2"].value).startswith("0001")
    assert sheet["A3"].value is None
    assert "1E2104" in str(sheet["N2"].value)
    assert "\n" not in str(sheet["N2"].value)


def test_path_groups():
    from src.pathways import MODE_AE, MODE_FULL_ST, MODE_SIM, MODE_ST2, extract_stuetzen, parse_path_nodes, route_groups

    nodes = parse_path_nodes("1E2104 (E 2,0) 1E2105 (E 1,0) 2EDV208 (E 0,0) 2EDV208.1 (E 2,3) 3E4005")
    assert nodes[0] == "1E2104"
    assert nodes[-1] == "3E4005"
    assert "(E" not in "".join(nodes)
    assert extract_stuetzen(nodes) == ["2EDV208", "2EDV208.1"]
    assert parse_path_nodes("KA (E 0,0) KE") == []
    assert parse_path_nodes("RingA (E 0,0) RingE") == []

    df = pd.DataFrame(
        {
            "join_key": ["a", "b", "c"],
            "Kabelnr": ["0002", "0003", "0004"],
            "Knoten_A": ["1E2104", "1E2104", "9X0001"],
            "Knoten_E": ["3E4005", "3E4005", "9X0002"],
            "Category": ["Cat3", "Cat1", "Cat1"],
            "Length_m": [128, 40, 10],
            "Pfad": [
                "1E2104 (E 2,0) 1E2105 (E 1,0) 2EDV208 (E 0,0) 3EDH206 (E 1,0) 3E2101 (E 3,1) 3E2102 (E 3,1) 3E2103 (E 3,1) 3E2104 (E 3,1) 3E2105 (E 3,1) 3E4005",
                "1E2104 (E 2,0) 1E2105 (E 1,0) 2EDV208 (E 0,0) 3EDH206 (E 1,0) 3E2101 (E 3,1) 3E2102 (E 3,1) 3E2103 (E 3,1) 3E2104 (E 3,1) 3E2105 (E 3,1) 3E4109",
                "KA (E 0,0) KE",
            ],
            "Pulled": [False, False, False],
            "verlegt": [pd.NaT, pd.NaT, pd.NaT],
            "Startsegment_1": [pd.NA, pd.NA, pd.NA],
            "Gestoppt_1": [pd.NaT, pd.NaT, pd.NaT],
        }
    )
    ae = route_groups(df, mode=MODE_AE, min_size=2)
    assert len(ae) == 1
    assert ae.iloc[0]["Point_A"] == "1E2104"
    assert ae.iloc[0]["Point_E"] == "3E4005"
    assert ae.iloc[0]["Cable_count"] == 2
    assert ae.iloc[0]["Cat3_n"] == 1
    assert ae.iloc[0]["Cat1_n"] == 1
    assert ae.iloc[0]["Total_m"] == 168

    st2 = route_groups(df, mode=MODE_ST2, min_size=2)
    assert len(st2) == 1
    assert "2EDV208" in st2.iloc[0]["Stuetzen"]
    assert "3EDH206" in st2.iloc[0]["Stuetzen"]

    full_st = route_groups(df, mode=MODE_FULL_ST, min_size=2)
    assert len(full_st) == 1
    assert full_st.iloc[0]["Cable_count"] == 2

    similar = route_groups(df, mode=MODE_SIM, min_size=2, similarity=0.9)
    assert len(similar) == 1
    assert similar.iloc[0]["Cable_count"] == 2
    assert "KA" not in similar.iloc[0]["Group"]
    assert "Ring" not in similar.iloc[0]["Group"]


def test_pull_length_calculator():
    from src.length_calc import (
        analyze_overlap,
        calculate_pulled_length,
        highlight_pulled_segment,
        highlight_route_layers,
        parse_route,
    )

    route = "3E2004 (E 3,1) 3E2005 (E 2,0) 1E3206 (E 1,5) 3EDV205"
    nodes, distances = parse_route(route)
    assert nodes == ["3E2004", "3E2005", "1E3206", "3EDV205"]
    assert distances == [3.1, 2.0, 1.5]

    result = calculate_pulled_length(route, "", 0.0, "1E3206", 0.5, "1E3206", "3EDV205")
    assert result.ok
    assert round(result.meters, 1) == 2.0

    reverse = calculate_pulled_length(route, "3E2004", 1.0, "", 0.0, "3EDV205", "3E2004")
    assert reverse.ok
    assert round(reverse.meters, 1) == 7.6

    missing = calculate_pulled_length(route, "", 0.0, "", 0.0, "NOPE", "3EDV205")
    assert not missing.ok

    html = highlight_pulled_segment(route, "1E3206", "3EDV205")
    assert "<mark" in html
    assert "1E3206" in html
    assert "3EDV205" in html
    assert "3E2004" in html
    assert html.index("<mark") < html.index("1E3206")
    assert "3E2004 (E 3,1)" in html or "3E2004" in html.split("<mark")[0]

    layered = highlight_route_layers(
        route,
        pulled_start="3E2004",
        pulled_end="1E3206",
        selected_start="3E2005",
        selected_end="3EDV205",
    )
    assert 'data-role="overlap"' in layered or "#f43f5e" in layered
    assert "#38bdf8" in layered or 'data-role="selected"' in layered

    advice = analyze_overlap(route, "3E2004", "1E3206", "3E2005", "3EDV205")
    assert advice.has_overlap
    assert advice.suggest_start == "1E3206"
    assert advice.suggest_end == "3EDV205"

    clean = analyze_overlap(route, "3E2004", "3E2005", "1E3206", "3EDV205")
    assert not clean.has_overlap


def test_cable_pull_status_not_pulled_is_not_complete():
    """Regression: 'Not pulled' contains substring 'pulled' — must not mean complete."""
    from src.analysis import PULL_NOT, PULL_PARTIAL, PULL_YES
    from src.length_calc import (
        STATUS_COMPLETE,
        STATUS_NONE,
        STATUS_PARTIAL,
        cable_pull_status,
        cable_status_message,
        has_verlegt_date,
    )

    # Like cable 3602: Pull_status text says Not pulled / no verlegt
    cable_3602 = pd.Series(
        {
            "Kabelnr": "3602",
            "verlegt": pd.NaT,
            "Pulled": False,
            "Pull_status": PULL_NOT,
            "Startsegment_1": pd.NA,
            "Endsegment_1": pd.NA,
        }
    )
    assert not has_verlegt_date(cable_3602)
    assert cable_pull_status(cable_3602) == STATUS_NONE
    assert "Teljesen" not in cable_status_message(cable_3602)
    assert "Nincs behúzás" in cable_status_message(cable_3602)

    for bad in ("Not pulled", "not pulled", "NOT PULLED"):
        row = pd.Series({"verlegt": pd.NaT, "Pulled": False, "Pull_status": bad})
        assert cable_pull_status(row) == STATUS_NONE

    partial = pd.Series(
        {
            "Kabelnr": "3602",
            "verlegt": pd.NaT,
            "Pulled": False,
            "Pull_status": PULL_PARTIAL,
            "Startsegment_1": "1E3206",
            "Endsegment_1": "3EDV205",
        }
    )
    assert cable_pull_status(partial) == STATUS_PARTIAL
    msg = cable_status_message(partial)
    assert "Megállítva" in msg
    assert "1E3206" in msg and "3EDV205" in msg

    done = pd.Series(
        {
            "Kabelnr": "3602",
            "verlegt": pd.Timestamp("2026-08-17"),
            "Pulled": True,
            "Pull_status": PULL_YES,
            "Startsegment_1": "1E3206",
            "Endsegment_1": "3EDV205",
        }
    )
    assert has_verlegt_date(done)
    assert cable_pull_status(done) == STATUS_COMPLETE

    # Pull_status == "Pulled" / Pulled=True but no verlegt → not complete
    fake_complete = pd.Series(
        {
            "Kabelnr": "3602",
            "verlegt": pd.NaT,
            "Pulled": True,
            "Pull_status": PULL_YES,
            "Startsegment_1": pd.NA,
            "Endsegment_1": pd.NA,
        }
    )
    assert cable_pull_status(fake_complete) == STATUS_NONE

    for empty in ("", "nein", "nan", None, pd.NaT):
        row = pd.Series(
            {
                "verlegt": empty,
                "Pulled": False,
                "Pull_status": PULL_NOT,
                "Startsegment_1": "",
                "Endsegment_1": "",
            }
        )
        assert cable_pull_status(row) == STATUS_NONE


def test_area_bahn_stuetzen_split():
    from src.areas import AreaDef, area_nodes_for_kind, areas_node_catalog

    area = AreaDef(
        key="room-a",
        name="Room A",
        deck="01",
        area_ids="A1",
        trays=["1E2104", "1EDV606", "1E2105"],
        penetrations=["3EDH407"],
    )
    assert area.bahnen == ["1E2104", "1E2105"]
    assert area.stuetzen == ["1EDV606", "3EDH407"]
    assert area_nodes_for_kind(area, "bahn") == ["1E2104", "1E2105"]
    assert area_nodes_for_kind(area, "stuetzen") == ["1EDV606", "3EDH407"]
    assert area_nodes_for_kind(area, "both") == ["1E2104", "1EDV606", "1E2105", "3EDH407"]
    catalog = areas_node_catalog([area])
    assert catalog[0]["bahn_count"] == 2
    assert catalog[0]["stuetzen_count"] == 2
    assert catalog[0]["all_count"] == 4


def test_stuetzen_bahn_filter():
    from openpyxl import load_workbook

    from src.export import catra_list_xlsx
    from src.stutzen_filter import criteria_header, filter_by_nodes, stutzen_filename

    catra = parse_catra_list(_catra_xlsx())
    and_hit = filter_by_nodes(catra, ["1E2104", "1E2105"], "AND")
    assert list(and_hit["join_key"]) == ["2"]

    or_hit = filter_by_nodes(catra, ["1EDH219", "1E2104"], "OR")
    assert set(or_hit["join_key"]) == {"2", "6", "7"}

    none = filter_by_nodes(catra, ["9EDV999"], "AND")
    assert none.empty

    matched = filter_by_nodes(catra, ["1E2104"], "OR")
    header = criteria_header(list(catra.attrs.get("catra_header") or []), ["1E2104"], "OR")
    payload = catra_list_xlsx(matched, header)
    assert len(matched) == 1
    assert stutzen_filename(["1E2104"]).startswith("stutzen_")
    sheet = load_workbook(BytesIO(payload)).active
    assert "Szűrési kritériumok" in str(sheet["A1"].value)
    assert str(sheet["A5"].value).startswith("0002")


def test_stutzen_filter_and_or():
    from src.stutzen_filter import filter_by_nodes, parse_search_nodes, search_nodes_stats, stutzen_filename

    df = pd.DataFrame(
        {
            "Kabelnr": ["1", "2", "3"],
            "Laenge": [10, 20, 30],
            "Pfad": [
                "1EDV606.1 (E 1,0) 1EDV606 (E 2,0) 1E2104",
                "1E1206 (E 3,0) 1EDH219",
                "1EDV606 (E 1,0) 1E2104",
            ],
        }
    )
    nodes = parse_search_nodes("1EDV606, 1E2104")
    both = filter_by_nodes(df, nodes, "AND")
    assert list(both["Kabelnr"]) == ["1", "3"]
    either = filter_by_nodes(df, nodes, "OR")
    assert list(either["Kabelnr"]) == ["1", "3"]
    dotted = filter_by_nodes(df, ["1EDV606.1"], "OR")
    assert list(dotted["Kabelnr"]) == ["1"]
    # case-insensitive Stütze / Bahn match
    lower = filter_by_nodes(df, ["1edv606", "1e2104"], "AND")
    assert list(lower["Kabelnr"]) == ["1", "3"]
    assert stutzen_filename(nodes) == "stutzen_1EDV606_1E2104.xlsx"

    unique_nodes = parse_search_nodes("1EDV606; 1E2104; 1EDV606", unique_only=True)
    assert unique_nodes == ["1EDV606", "1E2104"]
    stats = search_nodes_stats("1EDV606; 1E2104; 1EDV606; 1e2104")
    assert stats["raw_count"] == 4
    assert stats["unique_count"] == 2
    assert stats["duplicate_count"] == 2


def test_stutzen_pull_status_filter_and_export_columns():
    from src.analysis import PULL_NOT, PULL_PARTIAL, PULL_YES
    from src.stutzen_filter import (
        filter_by_pull_statuses,
        pull_status_counts,
        select_export_columns,
    )

    df = pd.DataFrame(
        {
            "Kabelnr": ["1", "2", "3"],
            "Liste": ["L1", "L2", "L3"],
            "Category": ["thin", "thick", "thin"],
            "Thickness": ["thin", "thick", "thin"],
            "Length_m": [10, 20, 30],
            "Pull_status": [PULL_NOT, PULL_PARTIAL, PULL_YES],
            "Geraet_A": ["A1", "A2", "A3"],
            "Geraet_E": ["E1", "E2", "E3"],
            "verlegt": [None, None, "12.08.2026"],
            "Gestoppt_1": [None, "01.09.2026", "01.08.2026"],
            "Startsegment_1": [None, "1E2104", "1E1000"],
            "Endsegment_1": [None, "1EDV606", "1E2000"],
            "Gestoppt_2": [None, "05.09.2026", None],
            "Startsegment_2": [None, "1EDV606", None],
            "Endsegment_2": [None, "1E3206", None],
            "Pfad": ["a", "b", "c"],
        }
    )
    counts = pull_status_counts(df)
    assert counts == {"not_started": 1, "started": 1, "fully_pulled": 1}
    from src.stutzen_filter import thickness_counts

    assert thickness_counts(df) == {"thick": 1, "thin": 2, "unknown": 0}
    started = filter_by_pull_statuses(df, ["started"])
    assert list(started["Kabelnr"]) == ["2"]
    slim = select_export_columns(started, ["kabelnr", "liste", "stop_location", "stop_locations"])
    assert list(slim.columns) == ["Kabelnr", "Liste", "Megállítás", "Megállítások"]
    # Only the overall current span (first start → latest end), not every stop.
    assert slim.iloc[0]["Megállítás"] == "1E2104 → 1E3206"
    assert slim.iloc[0]["Megállítások"] == "1E2104 → 1E3206"
    fully = filter_by_pull_statuses(df, ["fully_pulled"])
    fully_slim = select_export_columns(fully, ["kabelnr", "stop_location", "stop_locations"])
    # verlegt → Megállítás empty even when Start/End remain filled
    assert fully_slim.iloc[0]["Megállítás"] == ""
    assert fully_slim.iloc[0]["Megállítások"] == ""


def test_partial_length_calculator():
    from src.length_calc import calculate_pulled_length, format_meters

    route = "3E2004 (E 3,1) 3E2005 (E 2,0) 1E3206"
    ok = calculate_pulled_length(route, "3E2004", 1.0, "1E3206", 0.5, "3E2004", "1E3206")
    assert ok.ok
    assert round(ok.meters, 1) == 6.6
    assert format_meters(ok.meters) == "6,6 m"
    swapped = calculate_pulled_length(route, "3E2004", 0, "1E3206", 0, "1E3206", "3E2004")
    assert round(swapped.meters, 1) == 5.1
    missing = calculate_pulled_length(route, "", 0, "", 0, "NINCS", "1E3206")
    assert not missing.ok
    empty = calculate_pulled_length("", "", 0, "", 0, "", "")
    assert not empty.ok


def test_stichmass_from_stuetze_to_a():
    from src.length_calc import stichmass_from_stuetze_to_a

    # A ----3.1---- B ----2.0---- C ----1.5---- D(=node_a)
    route = "3E2004 (E 3,1) 3E2005 (E 2,0) 1E3206 (E 1,5) 3EDV205"
    # Stütze at B → A-end D: 2.0 + 1.5 + Aln_A 0.8 = 4.3
    ok = stichmass_from_stuetze_to_a(route, "3E2005", "3EDV205", 0.8)
    assert ok.ok
    assert round(ok.route_m, 1) == 3.5
    assert round(ok.allowance_a, 1) == 0.8
    assert round(ok.meters, 1) == 4.3

    # Stütze == A-end: only Aln_A
    at_a = stichmass_from_stuetze_to_a(route, "3EDV205", "3EDV205", 1.2)
    assert at_a.ok
    assert round(at_a.route_m, 1) == 0.0
    assert round(at_a.meters, 1) == 1.2

    missing = stichmass_from_stuetze_to_a(route, "NINCS", "3EDV205", 0.5)
    assert not missing.ok
    assert "nem található" in missing.error.lower()


def test_workforce_headcount_and_positions():
    from openpyxl import Workbook

    from src.workforce import (
        daily_headcount,
        daily_hours,
        headcount_growth,
        is_present,
        latest_headcount_day,
        meters_per_hour,
        meters_per_person,
        parse_workforce_excel,
        people_present_on,
        unique_positions,
    )

    assert is_present(10)
    assert is_present("9,5")
    assert not is_present("H")
    assert not is_present(0)
    assert not is_present(None)

    wb = Workbook()
    ws = wb.active
    ws.title = "Szeptember"
    ws.append([None, None, None])
    ws.append([None, "Name", "position", 1, 2, 3])
    ws.append([None, "Alice", "Cable puller", 10, "H", 10])
    ws.append([None, "Bob", "Electrician", 10, 10, 0])
    ws.append([None, "Sum", "Cable puller", 20, 10, 10])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    df = parse_workforce_excel(buf, year=2026)
    assert set(df["Name"]) == {"Alice", "Bob"}
    assert unique_positions(df) == ["Cable puller", "Electrician"]

    daily = daily_headcount(df)
    day1 = daily.loc[daily["Date"].eq(pd.Timestamp("2026-09-01"))].iloc[0]
    assert int(day1["Total"]) == 2
    assert int(day1["Cable puller"]) == 1
    assert int(day1["Electrician"]) == 1

    day2 = daily.loc[daily["Date"].eq(pd.Timestamp("2026-09-02"))].iloc[0]
    assert int(day2["Total"]) == 1  # Alice H, Bob present
    assert int(day2["Cable puller"]) == 0

    hours = daily_hours(df)
    assert float(hours.loc[hours["Date"].eq("2026-09-01"), "Hours"].iloc[0]) == 20.0
    assert float(hours.loc[hours["Date"].eq("2026-09-02"), "Hours"].iloc[0]) == 10.0
    assert float(hours.loc[hours["Date"].eq("2026-09-03"), "Hours"].iloc[0]) == 10.0
    assert int(hours.loc[hours["Date"].eq("2026-09-01"), "Headcount"].iloc[0]) == 2

    # Trailing empty days are trimmed; weekends can be dropped
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "Szeptember"
    ws2.append([None, None, None])
    # 2026-09-04 Friday, 05 Sat, 06 Sun, 07 Mon with people, then empty days in sheet to 10
    ws2.append([None, "Name", "position", 4, 5, 6, 7, 8, 9, 10])
    ws2.append([None, "Alice", "Cable puller", 10, 10, "H", 10, None, None, None])
    buf2 = BytesIO()
    wb2.save(buf2)
    buf2.seek(0)
    df2 = parse_workforce_excel(buf2, year=2026)
    full = daily_headcount(df2, exclude_weekends=False, trim_empty_tail=True, only_days_with_people=True)
    assert pd.Timestamp(full["Date"].max()) == pd.Timestamp("2026-09-07")
    assert int(full.loc[full["Date"].eq("2026-09-07"), "Total"].iloc[0]) == 1
    assert (full["Total"] > 0).all()
    # Sep 6 is Sunday with H only → not present → dropped
    assert pd.Timestamp("2026-09-06") not in set(pd.to_datetime(full["Date"]))
    no_we = daily_headcount(df2, exclude_weekends=True, trim_empty_tail=True, only_days_with_people=True)
    assert all(pd.to_datetime(no_we["Date"]).dt.dayofweek < 5)
    assert pd.Timestamp("2026-09-05") not in set(pd.to_datetime(no_we["Date"]))
    assert (no_we["Total"] > 0).all()

    only_pull = daily_headcount(df, ["Cable puller"])
    assert int(only_pull.loc[only_pull["Date"].eq(pd.Timestamp("2026-09-01")), "Total"].iloc[0]) == 1

    when, count = latest_headcount_day(daily, as_of=date(2026, 9, 3))
    assert when == pd.Timestamp("2026-09-03")
    assert count == 1  # Alice 10, Bob 0

    people = people_present_on(df, date(2026, 9, 1))
    assert list(people["Name"]) == ["Alice", "Bob"]

    growth = headcount_growth(daily)
    assert growth["start"] == 2
    assert growth["end"] == 1
    assert growth["peak"] == 2

    meters = pd.DataFrame(
        {
            "Date": pd.to_datetime(["2026-09-01", "2026-09-02", "2026-09-03"]),
            "Total": [20.0, 10.0, 5.0],
        }
    )
    prod = meters_per_person(daily, meters)
    assert round(float(prod.loc[prod["Date"].eq("2026-09-01"), "M_per_person"].iloc[0]), 1) == 10.0
    assert round(float(prod.loc[prod["Date"].eq("2026-09-02"), "M_per_person"].iloc[0]), 1) == 10.0
    assert float(prod.loc[prod["Date"].eq("2026-09-03"), "M_per_person"].iloc[0]) == 5.0

    per_h = meters_per_hour(hours, meters)
    assert round(float(per_h.loc[per_h["Date"].eq("2026-09-01"), "M_per_hour"].iloc[0]), 2) == 1.0
    assert round(float(per_h.loc[per_h["Date"].eq("2026-09-02"), "M_per_hour"].iloc[0]), 2) == 1.0
    assert round(float(per_h.loc[per_h["Date"].eq("2026-09-03"), "M_per_hour"].iloc[0]), 2) == 0.5
    assert float(per_h.loc[per_h["Date"].eq("2026-09-01"), "Hours"].iloc[0]) == 20.0
    assert float(per_h.loc[per_h["Date"].eq("2026-09-01"), "Meters"].iloc[0]) == 20.0


def _gesamt_xlsx_with_liste_separator() -> BytesIO:
    """New Kabelziehliste-style export with a Liste: section row between blocks."""
    wb = Workbook()
    ws = wb.active
    ws.append(["CaTra - Kabelziehliste"])
    ws.append(["Objekt: 598"])
    ws.append(["Erstellt: am 28.09.2026"])
    ws.append(["Kabel", "Artikel / Zeichnung", None, "Länge", "ALP", None, "Knoten", "A.län.", "Ort", "Gerät"])
    ws.append(["Liste: 0005 | Ringkabel"])
    ws.append(["5000 (E) K1", "Ship PWR", "2x25", 25, None, "A:", "1EDV217.1", 7, "1D", "Main Dist"])
    ws.append(["MDS", "1a MDS", None, None, None, "E:", "1E2119", 5.5, "1D", "24V Dist"])
    ws.append(["L0005 TKF"])
    ws.append(["1EDV217.1 (E 0,0) 1EDV217 (E 1,0) 1E2119"])
    ws.append(["Liste: 0008 | Other"])
    ws.append(["5001 (E) K1", "Ship PWR", "3x95", 40, None, "A:", "1E1206", 4, "1D", "DG"])
    ws.append(["PWR", "0008", None, None, None, "E:", "1EDV207.1", 4, "1D", "Converter"])
    ws.append(["L0008 Stolze"])
    ws.append(["1E1206 (E 2,0) 1EDV207.1"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_parse_catra_liste_separator_and_liste_field():
    catra = parse_catra_list(_gesamt_xlsx_with_liste_separator())
    assert set(catra["join_key"]) == {"5000", "5001"}
    assert list(catra["Liste"]) == ["L0005", "L0008"]
    assert "1EDV217.1" in str(catra.loc[catra["join_key"] == "5000", "Pfad"].iloc[0])


def test_reconcile_unmatched_and_length_conflict():
    from src.join import unmatched_ids
    from src.reconcile import enrich_joined_for_reconcile, field_mismatches, reconcile_report

    status = pd.DataFrame(
        {
            "Kabelnr": ["5000", "9999"],
            "join_key": ["5000", "9999"],
            "Zuschn": [30.0, 10.0],
            "A_Geraet": ["Main Dist", "X"],
            "E_Geraet": ["24V Dist", "Y"],
            "Liste": ["0005", "1"],
            "verlegt": [pd.NaT, pd.NaT],
            "Total_teilgezogene_meter": [0.0, 0.0],
            "Startsegment_1": [pd.NA, pd.NA],
            "Endsegment_1": [pd.NA, pd.NA],
            "Gestoppt_1": [pd.NaT, pd.NaT],
            "Pulled": [False, False],
        }
    )
    catra = parse_catra_list(_gesamt_xlsx_with_liste_separator())
    only_status, only_catra = unmatched_ids(status, catra)
    assert "9999" in only_status
    assert "5001" in only_catra

    report = reconcile_report(status, catra)
    assert "9999" in report["only_status"]
    assert "5001" in report["only_catra"]

    joined = join_cable_data(status, catra)
    joined = enrich_joined_for_reconcile(status, catra, joined)
    conflicts = field_mismatches(joined)
    # 5000: Zuschn 30 vs Laenge 25 → Length_m conflict
    length_rows = conflicts[conflicts["field"] == "Length_m"]
    assert not length_rows.empty
    assert str(length_rows.iloc[0]["Kabelnr"]) == "5000"
    assert float(length_rows.iloc[0]["progress_value"]) == 30.0
    assert float(length_rows.iloc[0]["gesamt_value"]) == 25.0


def test_areas_membership_remaining_hours_eta_blocking():
    from datetime import date

    from src.areas import (
        AreaDef,
        affected_areas,
        classify_cable_in_area,
        estimate_schedule,
        parse_areas_excel,
        parse_node_list,
        person_hours,
        summarize_all_areas,
        summarize_area,
    )

    assert parse_node_list("1E5101; 1E5103;1E5102;") == ["1E5101", "1E5103", "1E5102"]

    hely_path = r"c:\Users\User\Desktop\Helység adatok.xlsx"
    try:
        areas_real = parse_areas_excel(hely_path)
        assert len(areas_real) >= 10
        assert any(a.name == "Misc. Room" for a in areas_real)
    except FileNotFoundError:
        pass

    area = AreaDef(
        key="Room|Main|1",
        name="Room",
        deck="Main",
        area_ids="1",
        trays=["1E2002", "1E2003"],
        penetrations=["1EDH201"],
    )
    route = "1E2001 (E 10,0) 1E2002 (E 20,0) 1EDH201 (E 30,0) 1E2004"

    row_none = pd.Series(
        {
            "Kabelnr": "100",
            "Pfad": route,
            "Length_m": 99.0,
            "Thickness": "Thin",
            "Knoten_A": "1E2001",
            "Knoten_E": "1E2004",
            "verlegt": pd.NaT,
            "Startsegment_1": None,
            "Endsegment_1": None,
            "Gestoppt_1": None,
        }
    )
    info = classify_cable_in_area(row_none, area)
    assert info is not None
    assert info["through"] is True
    assert info["remaining_m"] == 99.0

    row_partial = row_none.copy()
    row_partial["Kabelnr"] = "101"
    row_partial["Startsegment_1"] = "1E2001"
    row_partial["Endsegment_1"] = "1E2002"
    row_partial["Gestoppt_1"] = date(2026, 9, 1)
    info_p = classify_cable_in_area(row_partial, area)
    assert info_p is not None
    assert info_p["stop"] is True
    assert round(info_p["remaining_m"], 1) == 20.0

    row_past = row_partial.copy()
    row_past["Kabelnr"] = "102"
    row_past["Endsegment_1"] = "1E2004"
    assert classify_cable_in_area(row_past, area)["remaining_m"] == 0.0

    row_done = row_none.copy()
    row_done["Kabelnr"] = "104"
    row_done["verlegt"] = date(2026, 9, 10)
    assert classify_cable_in_area(row_done, area)["remaining_m"] == 0.0

    hours = person_hours(40.0, 6.0, thin_pace=20.0, thick_pace=3.0, unchecked_stuetzen=2, h_per_stuetze=5.0)
    assert round(hours, 1) == 14.0

    sched = estimate_schedule(112.0, pullers=2, hours_per_week=56.0, as_of=date(2026, 9, 7))
    assert sched["workdays"] == 5
    assert sched["end_date"] == date(2026, 9, 11)
    assert sched["kw"] == "KW37 2026"

    zero = estimate_schedule(0.0, pullers=2, hours_per_week=56.0, as_of=date(2026, 9, 7))
    assert zero["workdays"] == 0
    assert zero["kw"] == "KW37 2026"

    df = pd.DataFrame(
        [
            {
                "Kabelnr": "200",
                "Pfad": route,
                "Length_m": 50.0,
                "Thickness": "Thin",
                "Knoten_A": "1E2001",
                "Knoten_E": "1E2004",
                "verlegt": date(2026, 9, 10),
                "Startsegment_1": "1E2001",
                "Endsegment_1": "1E2004",
                "Gestoppt_1": date(2026, 9, 1),
            },
            {
                "Kabelnr": "201",
                "Pfad": route,
                "Length_m": 80.0,
                "Thickness": "Thick",
                "Knoten_A": "1E2001",
                "Knoten_E": "1E2004",
                "verlegt": pd.NaT,
                "Startsegment_1": None,
                "Endsegment_1": None,
                "Gestoppt_1": None,
            },
        ]
    )
    summary = summarize_area(
        area, df, installed_stuetzen=[], thin_pace=20, thick_pace=3, h_per_stuetze=5, pullers=1, hours_per_week=56
    )
    assert summary["cables"] == 2
    assert summary["pulled"] == 1
    assert round(summary["remaining_thick_m"], 1) == 80.0

    area_b = AreaDef(
        key="Hall|Main|2",
        name="Hall",
        deck="Main",
        area_ids="2",
        trays=["1E2004"],
        penetrations=[],
    )
    summaries = summarize_all_areas([area, area_b], df)
    affected = affected_areas(f"area:{area.key}", [area, area_b], summaries, df)
    assert any(a.key == area_b.key for a in affected)

    from src.areas import area_cables_frame, remaining_cables_frame

    all_cables = area_cables_frame(summary, df)
    assert set(all_cables["Kabelnr"].astype(str)) == {"200", "201"}
    rem = remaining_cables_frame(summary, df)
    assert list(rem["Kabelnr"].astype(str)) == ["201"]


def test_summarize_all_areas_single_pass_matches_per_area():
    """One-pass aggregation must match classic per-area summarize_area results."""
    from src.areas import AreaDef, summarize_all_areas, summarize_area

    route = "1E2001 (E 2,0) 1EDV100 (E 3,0) 1E2004"
    area = AreaDef(
        key="Room|1|A",
        name="Room",
        deck="1",
        area_ids="A",
        trays=["1E2001", "1EDV100", "1E2004"],
        penetrations=[],
    )
    area_b = AreaDef(
        key="Hall|Main|2",
        name="Hall",
        deck="Main",
        area_ids="2",
        trays=["1E2004"],
        penetrations=[],
    )
    df = pd.DataFrame(
        [
            {
                "Kabelnr": "200",
                "Pfad": route,
                "Length_m": 50.0,
                "Thickness": "Thin",
                "Knoten_A": "1E2001",
                "Knoten_E": "1E2004",
                "verlegt": date(2026, 9, 10),
                "Startsegment_1": "1E2001",
                "Endsegment_1": "1E2004",
                "Gestoppt_1": date(2026, 9, 1),
            },
            {
                "Kabelnr": "201",
                "Pfad": route,
                "Length_m": 80.0,
                "Thickness": "Thick",
                "Knoten_A": "1E2001",
                "Knoten_E": "1E2004",
                "verlegt": pd.NaT,
                "Startsegment_1": None,
                "Endsegment_1": None,
                "Gestoppt_1": None,
            },
        ]
    )
    one_pass = {
        s["area"].key: s
        for s in summarize_all_areas([area, area_b], df, apply_schedule=False)
    }
    for a in (area, area_b):
        classic = summarize_area(a, df)
        fast = one_pass[a.key]
        assert classic["cables"] == fast["cables"]
        assert classic["pulled"] == fast["pulled"]
        assert round(classic["remaining_m"], 3) == round(fast["remaining_m"], 3)
        assert round(classic["person_hours"], 3) == round(fast["person_hours"], 3)
        assert set(classic["cable_ids"]) == set(fast["cable_ids"])

