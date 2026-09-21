from __future__ import annotations

import math
import re
import zipfile
from datetime import datetime, date
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import TYPE_STRING
from openpyxl.styles import Alignment, Font

_PATH_TOKEN = re.compile(r"\S+(?:\s+\([^)]*\))?")
_PATH_WIDTH = 88
_PATH_ALIGN = Alignment(wrap_text=True, vertical="top")

_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_MAX_CELL = 32767

DEFAULT_CATRA_HEADER = [
    ["CaTra - Kabelliste"],
    ["Objekt"],
    ["Erstellt"],
    ["Kabel", "Artikel / Liste", None, "Länge", "ALP", None, "Knoten", "A.län.", "Ort", "Gerät"],
]


def to_xlsx_bytes(df: pd.DataFrame, sheet_name: str = "Data") -> bytes:
    buffer = BytesIO()
    export = df.copy()
    drop = [col for col in export.columns if col.startswith("_")]
    if drop:
        export = export.drop(columns=drop)
    for col in export.columns:
        if pd.api.types.is_datetime64_any_dtype(export[col]):
            export[col] = export[col].dt.strftime("%d.%m.%Y")
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        export.to_excel(writer, index=False, sheet_name=sheet_name[:31])
    return buffer.getvalue()


def safe_filename(name: str) -> str:
    cleaned = re.sub(r"[^\w.\-]+", "_", str(name).strip())
    return cleaned[:80] or "location"


def _safe_excel_value(value):
    if value is None or value is pd.NA:
        return None
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value) or pd.isna(value):
            return None
        return value
    if hasattr(value, "item") and not isinstance(value, (bytes, str, list, tuple, dict)):
        try:
            value = value.item()
        except Exception:
            value = str(value)
        return _safe_excel_value(value)
    if isinstance(value, str):
        cleaned = _ILLEGAL_XML.sub("", value).replace("\x00", "")
        if len(cleaned) > _MAX_CELL:
            cleaned = cleaned[:_MAX_CELL]
        return cleaned or None
    if isinstance(value, (list, tuple, dict)):
        return _safe_excel_value(str(value))
    return value


def wrap_pathway(text, width: int = _PATH_WIDTH) -> str:
    raw = "" if text is None else str(text).strip()
    if not raw:
        return ""
    parts = _PATH_TOKEN.findall(raw.replace("\n", " "))
    if not parts:
        return raw
    lines = []
    current = parts[0]
    for part in parts[1:]:
        trial = f"{current} {part}"
        if len(trial) <= width:
            current = trial
        else:
            lines.append(current)
            current = part
    lines.append(current)
    return "\n".join(lines)


def _path_text(row: list) -> str:
    return " ".join(str(value).strip() for value in row if value is not None and str(value).strip())


def write_path_row(ws, row: int, text) -> None:
    wrapped = wrap_pathway(text)
    write_cell(ws, row, 1, wrapped)
    cell = ws.cell(row=row, column=1)
    cell.alignment = _PATH_ALIGN
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
    lines = wrapped.count("\n") + 1 if wrapped else 1
    ws.row_dimensions[row].height = max(18, min(15 * lines + 8, 150))


def write_cell(ws, row: int, col: int, value) -> None:
    """Write a value as data, never as an Excel formula."""
    value = _safe_excel_value(value)
    if value is None or value == "":
        return
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        cell.value = value
        cell.data_type = TYPE_STRING
        return
    cell.value = value


def catra_block_for_row(row: pd.Series) -> list[list]:
    block = row.get("_block")
    if isinstance(block, list) and block:
        return block
    return [
        [
            row.get("Kabel_raw") or row.get("Kabelnr"),
            row.get("Artikel") or row.get("Kabeltyp"),
            row.get("Typ") or row.get("Kabeladern"),
            row.get("Laenge") if pd.notna(row.get("Laenge")) else row.get("Length_m"),
            None,
            "A:",
            row.get("Knoten_A"),
            row.get("Aln_A"),
            row.get("Ort_A"),
            row.get("Geraet_A") or row.get("A_Geraet"),
        ],
        [
            row.get("Kabelbez"),
            row.get("Liste"),
            None,
            None,
            None,
            "E:",
            row.get("Knoten_E"),
            row.get("Aln_E"),
            row.get("Ort_E"),
            row.get("Geraet_E") or row.get("E_Geraet"),
        ],
        [row.get("Notizen") or row.get("Bemerkung"), None, None, None, None, "Z:", row.get("Knoten_Z")],
        [row.get("Pfad")],
    ]


def cables_at_location(joined: pd.DataFrame, detail: pd.DataFrame, location: str) -> pd.DataFrame:
    keys = set(detail.loc[detail["Stop_location"].eq(location), "join_key"].dropna())
    if not keys:
        nrs = set(detail.loc[detail["Stop_location"].eq(location), "Kabelnr"].dropna().astype(str))
        return joined[joined["Kabelnr"].astype(str).isin(nrs)]
    return joined[joined["join_key"].isin(keys)]


def catra_list_xlsx(cables: pd.DataFrame, header_rows: list[list] | None = None) -> bytes:
    header = header_rows or DEFAULT_CATRA_HEADER
    wb = Workbook()
    ws = wb.active
    ws.title = "Ergebnis"
    for r_idx, row in enumerate(header, start=1):
        for c_idx, value in enumerate(row, start=1):
            write_cell(ws, r_idx, c_idx, value)
        if r_idx == 1:
            ws.cell(1, 1).font = Font(bold=True, size=14)
        if r_idx == 4:
            for cell in ws[4]:
                cell.font = Font(bold=True)

    start = len(header) + 1
    for i, (_, cable) in enumerate(cables.iterrows()):
        block = catra_block_for_row(cable)
        for br, row in enumerate(block):
            excel_row = start + i * 4 + br
            if br == 3:
                write_path_row(ws, excel_row, _path_text(row) or cable.get("Pfad"))
                continue
            for c_idx, value in enumerate(row, start=1):
                write_cell(ws, excel_row, c_idx, value)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["J"].width = 36
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _first_value(row: pd.Series, *names):
    for name in names:
        if name not in row.index:
            continue
        value = row.get(name)
        if value is None or value == "":
            continue
        try:
            if pd.isna(value):
                continue
        except (TypeError, ValueError):
            pass
        return value
    return None


def _plain_path(row: pd.Series) -> str:
    block = row.get("_block")
    if isinstance(block, list) and len(block) >= 4:
        text = _path_text(block[3])
        if text:
            return " ".join(text.split())
    pfad = row.get("Pfad")
    if pfad is None:
        return ""
    return " ".join(str(pfad).split())


def kurzliste_frame(cables: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, cable in cables.iterrows():
        verlegt = _first_value(cable, "verlegt")
        if isinstance(verlegt, (pd.Timestamp, datetime)) and pd.notna(verlegt):
            verlegt = pd.Timestamp(verlegt).strftime("%d.%m.%Y")
        elif verlegt is not None:
            try:
                if pd.isna(verlegt):
                    verlegt = None
            except (TypeError, ValueError):
                pass
        rows.append(
            {
                "Kabelnr": _first_value(cable, "Kabel_raw", "Kabelnr"),
                "Liste": _first_value(cable, "Liste"),
                "Artikel / Typ": _first_value(cable, "Artikel", "Kabeltyp"),
                "Adern": _first_value(cable, "Typ", "Kabeladern"),
                "Länge": _first_value(cable, "Laenge", "Length_m"),
                "Category": _first_value(cable, "Category"),
                "Pull_status": _first_value(cable, "Pull_status"),
                "Knoten_A": _first_value(cable, "Point_A", "Knoten_A"),
                "Gerät_A": _first_value(cable, "Geraet_A", "A_Geraet"),
                "Ort_A": _first_value(cable, "Ort_A"),
                "Knoten_E": _first_value(cable, "Point_E", "Knoten_E"),
                "Gerät_E": _first_value(cable, "Geraet_E", "E_Geraet"),
                "Ort_E": _first_value(cable, "Ort_E"),
                "Pfad": _plain_path(cable),
                "Remaining_m": _first_value(cable, "Remaining_m"),
                "verlegt": verlegt,
            }
        )
    return pd.DataFrame(rows, columns=KURZ_COLUMNS)


KURZ_COLUMNS = [
    "Kabelnr",
    "Liste",
    "Artikel / Typ",
    "Adern",
    "Länge",
    "Category",
    "Pull_status",
    "Knoten_A",
    "Gerät_A",
    "Ort_A",
    "Knoten_E",
    "Gerät_E",
    "Ort_E",
    "Pfad",
    "Remaining_m",
    "verlegt",
]


def kurzliste_xlsx(cables: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kurzliste"
    frame = kurzliste_frame(cables)
    headers = KURZ_COLUMNS if frame.empty else list(frame.columns)
    for c_idx, name in enumerate(headers, start=1):
        write_cell(ws, 1, c_idx, name)
        ws.cell(1, c_idx).font = Font(bold=True)
    for r_idx, row in enumerate(frame.itertuples(index=False), start=2):
        for c_idx, value in enumerate(row, start=1):
            write_cell(ws, r_idx, c_idx, value)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["N"].width = 80
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).coordinate}"
    ws.freeze_panes = "A2"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def catra_locations_zip(
    joined: pd.DataFrame,
    detail: pd.DataFrame,
    locations: list[str],
    header_rows: list[list] | None = None,
) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for location in locations:
            cables = cables_at_location(joined, detail, location)
            if cables.empty:
                continue
            stem = safe_filename(location)
            archive.writestr(f"huzokartya/stopped_{stem}.xlsx", catra_list_xlsx(cables, header_rows))
            archive.writestr(f"kurzliste/stopped_{stem}.xlsx", kurzliste_xlsx(cables))
    return buffer.getvalue()

